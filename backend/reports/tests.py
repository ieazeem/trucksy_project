from datetime import timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from fleet.models import Driver
from loads.models import Load
from tracking.models import LocationPing

from .excel import haversine_km, parse_period, route_km

User = get_user_model()
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ReportTests(APITestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_demo", "--reset", verbosity=0)
        cls.admin = User.objects.get(username="admin")
        cls.yasir = Driver.objects.get(phone="34180130")

    def setUp(self):
        self.client.force_authenticate(self.admin)

    def wb(self, url):
        r = self.client.get(url)
        self.assertEqual(r.status_code, 200, getattr(r, "content", b"")[:300])
        self.assertEqual(r["Content-Type"], XLSX)
        self.assertIn("attachment;", r["Content-Disposition"])
        return load_workbook(BytesIO(r.content))

    def test_exports_are_dispatcher_only(self):
        self.client.force_authenticate(None)
        for url in ["/api/reports/operations.xlsx", "/api/reports/loads.xlsx", "/api/reports/drivers.xlsx",
                    f"/api/reports/drivers/{self.yasir.id}/route.xlsx", f"/api/drivers/{self.yasir.id}/history/"]:
            self.assertEqual(self.client.get(url).status_code, 401, url)
        self.client.force_authenticate(self.yasir.user)  # a driver, not staff
        self.assertEqual(self.client.get("/api/reports/operations.xlsx").status_code, 403)

    def test_operations_report_has_all_tabs_and_matching_numbers(self):
        today = timezone.localdate()
        wb = self.wb(f"/api/reports/operations.xlsx?from={today - timedelta(days=6)}&to={today}")
        self.assertEqual(wb.sheetnames, ["Summary", "Daily", "Loads", "Drivers", "Fleet"])
        delivered = Load.objects.filter(status="delivered",
                                        delivered_at__gte=parse_period({"from": str(today - timedelta(days=6)),
                                                                        "to": str(today)})[2]).count()
        summary = {wb["Summary"].cell(r, 1).value: wb["Summary"].cell(r, 2).value for r in range(4, 12)}
        self.assertEqual(summary["Delivered"], delivered)
        daily = wb["Daily"]
        self.assertEqual(daily.max_row, 1 + 7 + 1)  # header + 7 days + total
        self.assertEqual(sum(daily.cell(r, 2).value for r in range(2, 9)), delivered)
        self.assertTrue(str(daily.cell(9, 2).value).startswith("=SUM("))  # live Excel formula

    def test_loads_export_matches_dashboard_filters(self):
        today = timezone.localdate()
        q = f"from={today - timedelta(days=30)}&to={today}&status=delivered"
        api_rows = self.client.get(f"/api/loads/?{q}").json()
        ws = self.wb(f"/api/reports/loads.xlsx?{q}")["Loads"]
        refs = [ws.cell(r, 1).value for r in range(4, ws.max_row)]  # title row, blank, header, ..., total
        self.assertEqual(sorted(refs), sorted(l["reference"] for l in api_rows))
        self.assertTrue(all(ws.cell(r, 2).value == "Delivered" for r in range(4, ws.max_row)))
        self.assertEqual(ws.cell(4, 9).number_format, '#,##0.000 "BHD"')

    def test_drivers_export_lists_everyone(self):
        ws = self.wb("/api/reports/drivers.xlsx")["Drivers"]
        names = {ws.cell(r, 1).value for r in range(4, ws.max_row)}
        self.assertEqual(names, set(Driver.objects.values_list("full_name", flat=True)))

    def test_route_history_by_date_range(self):
        today = timezone.localdate()
        r = self.client.get(f"/api/drivers/{self.yasir.id}/history/?from={today}&to={today}")
        self.assertEqual(r.status_code, 200)
        data = r.json()
        self.assertGreater(data["stats"]["points"], 0)
        self.assertGreater(data["stats"]["km"], 0)
        times = [p["at"] for p in data["points"]]
        self.assertEqual(times, sorted(times))
        # A range with no driving returns an empty, valid result.
        old = today - timedelta(days=200)
        empty = self.client.get(f"/api/drivers/{self.yasir.id}/history/?from={old}&to={old}").json()
        self.assertEqual(empty["stats"]["points"], 0)
        self.assertIsNone(empty["stats"]["first"])
        # Excel version
        wb = self.wb(f"/api/reports/drivers/{self.yasir.id}/route.xlsx?from={today}&to={today}")
        self.assertEqual(wb.sheetnames, ["Summary", "Deliveries", "GPS points"])
        gps = wb["GPS points"]
        self.assertEqual(gps.max_row - 1, data["stats"]["points"])
        self.assertIn("google.com/maps", gps.cell(2, 6).hyperlink.target)

    def test_period_parsing_is_safe(self):
        d_from, d_to, start, end = parse_period({"from": "garbage", "to": "2026-10-05"}, default_days=7)
        self.assertEqual(str(d_to), "2026-10-05")
        self.assertEqual((d_to - d_from).days, 6)
        self.assertEqual((end - start).days, 7)
        a, b, *_ = parse_period({"from": "2026-10-09", "to": "2026-10-01"})  # swapped -> fixed
        self.assertLess(a, b)

    def test_distance_math(self):
        # Manama (Bab Al Bahrain) to Riffa is roughly 12 km as the crow flies.
        self.assertAlmostEqual(haversine_km((26.2361, 50.5770), (26.1300, 50.5550)), 12.0, delta=1.5)
        # A GPS glitch jump (> 3 km between two pings) is ignored.
        self.assertLess(route_km([(26.20, 50.58), (26.201, 50.581), (25.0, 51.0), (26.202, 50.582)]), 1)
