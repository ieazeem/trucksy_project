"""Small helpers so every Trucksy Excel file looks the same: sign-blue header row,
frozen header, filters on, sensible widths, BHD money format, Bahrain local times."""
from datetime import date, datetime, timedelta
from io import BytesIO
from math import asin, cos, radians, sin, sqrt

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SIGN_BLUE = "12459B"
AMBER = "F2A900"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=SIGN_BLUE)
TOTAL_FONT = Font(bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="FDF3D6")
THIN = Side(style="thin", color="D9DDD5")

MONEY = '#,##0.000 "BHD"'
KG = '#,##0 "kg"'
KM = '#,##0.0 "km"'
DT = "dd mmm yyyy hh:mm"
D = "dd mmm yyyy"


def local(dt):
    """Aware datetime -> naive Bahrain time (Excel has no time zones)."""
    if not dt:
        return None
    return timezone.localtime(dt).replace(tzinfo=None)


def parse_period(params, default_days=7):
    """?from=YYYY-MM-DD&to=YYYY-MM-DD (inclusive, Bahrain dates). Defaults to the last N days.
    Returns (from_date, to_date, start_aware, end_aware_exclusive)."""
    today = timezone.localdate()

    def _d(key):
        try:
            return date.fromisoformat(params.get(key, ""))
        except ValueError:
            return None

    d_to = _d("to") or today
    d_from = _d("from") or (d_to - timedelta(days=default_days - 1))
    if d_from > d_to:
        d_from, d_to = d_to, d_from
    tz = timezone.get_current_timezone()
    start = timezone.make_aware(datetime.combine(d_from, datetime.min.time()), tz)
    end = timezone.make_aware(datetime.combine(d_to + timedelta(days=1), datetime.min.time()), tz)
    return d_from, d_to, start, end


def haversine_km(a, b):
    lat1, lng1, lat2, lng2 = map(radians, (a[0], a[1], b[0], b[1]))
    h = sin((lat2 - lat1) / 2) ** 2 + cos(lat1) * cos(lat2) * sin((lng2 - lng1) / 2) ** 2
    return 2 * 6371.0 * asin(sqrt(h))


def route_km(points):
    """Distance along GPS points, ignoring impossible jumps (GPS glitches > 3 km between pings)."""
    km = 0.0
    for a, b in zip(points, points[1:]):
        step = haversine_km(a, b)
        if step < 3:
            km += step
    return km


class Sheet:
    """Write a formatted table: Sheet(wb, "Loads", [("Reference", 14), ("Price", 12, MONEY)])."""

    def __init__(self, wb, title, columns, first=False, title_row=None):
        self.ws = wb.active if first else wb.create_sheet()
        self.ws.title = title[:31]
        self.columns = columns
        self.start_row = 1
        if title_row:
            self.ws.cell(row=1, column=1, value=title_row).font = Font(bold=True, size=13, color=SIGN_BLUE)
            self.start_row = 3
        for i, col in enumerate(columns, start=1):
            c = self.ws.cell(row=self.start_row, column=i, value=col[0])
            c.font, c.fill = HEADER_FONT, HEADER_FILL
            c.alignment = Alignment(vertical="center")
            self.ws.column_dimensions[get_column_letter(i)].width = col[1]
        self.ws.freeze_panes = self.ws.cell(row=self.start_row + 1, column=1)
        self.row = self.start_row

    def add(self, values):
        self.row += 1
        for i, (v, col) in enumerate(zip(values, self.columns), start=1):
            c = self.ws.cell(row=self.row, column=i, value=v)
            if len(col) > 2 and v is not None:
                c.number_format = col[2]
            c.border = Border(bottom=THIN)

    def finish(self, totals=None):
        """totals: {column_index(1-based): "sum" | "count" | text}"""
        last_data = self.row
        if self.row > self.start_row:
            self.ws.auto_filter.ref = f"A{self.start_row}:{get_column_letter(len(self.columns))}{last_data}"
        if totals:
            self.row += 1
            for idx, kind in totals.items():
                letter = get_column_letter(idx)
                rng = f"{letter}{self.start_row + 1}:{letter}{max(last_data, self.start_row + 1)}"
                value = {"sum": f"=SUM({rng})", "count": f"=COUNTA({rng})"}.get(kind, kind)
                c = self.ws.cell(row=self.row, column=idx, value=value)
                c.font, c.fill = TOTAL_FONT, TOTAL_FILL
                col = self.columns[idx - 1]
                if kind == "sum" and len(col) > 2:
                    c.number_format = col[2]
        return self


def xlsx_response(wb: Workbook, filename: str) -> HttpResponse:
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp
