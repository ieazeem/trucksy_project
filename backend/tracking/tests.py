from datetime import timedelta

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.utils import timezone
from rest_framework.test import APITestCase

from fleet.models import Driver
from loads.models import Load
from tracking.models import LocationPing

User = get_user_model()


class SignalGapTests(APITestCase):
    """A phone that is switched off during a job must show up: live alert + gap in the route history."""

    @classmethod
    def setUpTestData(cls):
        call_command("seed_demo", "--reset", verbosity=0)
        cls.admin = User.objects.get(username="admin")

    def setUp(self):
        self.client.force_authenticate(self.admin)
        self.driver = Driver.objects.get(phone="34180130")
        Load.objects.filter(driver=self.driver, status__in=Load.ACTIVE_STATUSES).update(status=Load.Status.DELIVERED)
        self.load = Load.objects.create(customer_name="Gap Test", pickup_address="Sitra", dropoff_address="Riffa",
                                        driver=self.driver, status=Load.Status.IN_TRANSIT)
        LocationPing.objects.filter(driver=self.driver).delete()

    def ping(self, minutes_ago, lat, lng):
        LocationPing.objects.create(driver=self.driver, load=self.load, lat=lat, lng=lng,
                                    recorded_at=timezone.now() - timedelta(minutes=minutes_ago))

    def test_phone_off_mid_job_is_a_gap_with_distance_moved(self):
        self.ping(60, 26.15, 50.62)
        self.ping(59, 26.151, 50.62)
        self.ping(25, 26.05, 50.55)   # 34 minutes of nothing, ~13 km further on
        self.ping(24, 26.049, 50.55)
        self.driver.last_seen = timezone.now() - timedelta(minutes=1)
        self.driver.save()
        today = timezone.localdate()
        h = self.client.get(f"/api/drivers/{self.driver.id}/history/?from={today}&to={today}").json()
        gaps = [g for g in h["gaps"] if not g["ongoing"]]
        self.assertEqual(len(gaps), 1)
        self.assertEqual(gaps[0]["minutes"], 34)
        self.assertGreater(gaps[0]["km_moved"], 10)
        self.assertEqual(gaps[0]["load"], self.load.reference)

    def test_silent_driver_on_job_raises_live_alert(self):
        self.ping(30, 26.15, 50.62)
        self.driver.last_seen = timezone.now() - timedelta(minutes=30)
        self.driver.save()
        alerts = self.client.get("/api/dashboard/").json()["alerts"]
        mine = [a for a in alerts if a["driver_id"] == self.driver.id]
        self.assertEqual(len(mine), 1)
        self.assertEqual(mine[0]["load"], self.load.reference)
        self.assertGreaterEqual(mine[0]["minutes"], 29)

    def test_no_alert_while_driver_reports(self):
        self.driver.last_seen = timezone.now() - timedelta(minutes=2)
        self.driver.save()
        alerts = self.client.get("/api/dashboard/").json()["alerts"]
        self.assertFalse([a for a in alerts if a["driver_id"] == self.driver.id])

    def test_route_excel_has_signal_gaps_sheet(self):
        self.ping(60, 26.15, 50.62)
        self.ping(20, 26.05, 50.55)
        today = timezone.localdate()
        from io import BytesIO
        from openpyxl import load_workbook
        r = self.client.get(f"/api/reports/drivers/{self.driver.id}/route.xlsx?from={today}&to={today}")
        wb = load_workbook(BytesIO(r.content))
        self.assertIn("Signal gaps", wb.sheetnames)
        self.assertEqual(wb["Signal gaps"].cell(2, 4).value, 40)
